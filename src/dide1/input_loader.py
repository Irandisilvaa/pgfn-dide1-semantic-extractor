from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterable
import re
from openpyxl import load_workbook
from .text_utils import normalizar

ALIASES_CABECALHOS = {
    "processo": {"processo", "numero_processo", "n_processo"},
    "classe": {"classe", "classe_judicial"},
    "origem": {"origem"},
    "orgao_julgador": {"orgao_julgador", "orgao julgador"},
    "data": {"data", "data_expediente"},
    "polo_ativo": {"polo_ativo", "polo ativo"},
    "polo_passivo": {"polo_passivo", "polo passivo"},
    "expediente": {"expediente"},
    "prazo": {"prazo"},
    "assunto": {"assunto"},
    "secao_subsecao": {"secao_subsecao", "secao/subsecao", "seção/subseção"},
    "link": {"link"},
    "decisao": {"decisao_judicial", "decisão judicial", "decisao", "decisão"},
    "nucleo": {"nucleo", "núcleo"},
}


def _normalizar_cabecalho(v: Any) -> str:
    s = normalizar("" if v is None else str(v))
    s = re.sub(r"[^a-z0-9/]+", "_", s).strip("_")
    return s


def _localizar_colunas(cabecalhos):
    cabecalhos_normalizados = {_normalizar_cabecalho(h): i for i, h in enumerate(cabecalhos)}
    saida = {}
    for chave, valores in ALIASES_CABECALHOS.items():
        candidatos = {_normalizar_cabecalho(v) for v in valores}
        for c in candidatos:
            if c in cabecalhos_normalizados:
                saida[chave] = cabecalhos_normalizados[c]
                break
    faltantes = [x for x in ("processo", "classe", "orgao_julgador", "expediente", "decisao") if x not in saida]
    if faltantes:
        raise ValueError(f"Colunas obrigatórias não encontradas: {faltantes}. Cabeçalhos: {cabecalhos}")
    return saida

@dataclass
class RegistroDecisao:
    linha_origem: int
    dados: dict


def iterar_registros(caminho: str, aba: str | None = None) -> Iterable[RegistroDecisao]:
    p = Path(caminho)
    if not p.exists():
        raise FileNotFoundError(caminho)
    pasta_trabalho = load_workbook(p, read_only=True, data_only=True)
    planilha = pasta_trabalho[aba] if aba else pasta_trabalho[pasta_trabalho.sheetnames[0]]
    linhas = planilha.iter_rows(values_only=True)
    cabecalhos = list(next(linhas))
    mapeamento = _localizar_colunas(cabecalhos)
    for linha_excel, linha in enumerate(linhas, start=2):
        d = {}
        for chave, indice in mapeamento.items():
            d[chave] = linha[indice] if indice < len(linha) else None
        d["decisao"] = "" if d.get("decisao") is None else str(d.get("decisao"))
        d["processo"] = "" if d.get("processo") is None else str(d.get("processo"))
        yield RegistroDecisao(linha_excel, d)
