from __future__ import annotations

from collections import defaultdict
from pathlib import Path
import csv
import json

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .constants import CATEGORIAS

PREENCHIMENTO_CABECALHO = "1F4E78"
FONTE_CABECALHO = "FFFFFF"
PREENCHIMENTO_CONTEXTO = "D9EAF7"
PREENCHIMENTO_CANDIDATO = "E2F0D9"
PREENCHIMENTO_HUMANO = "FFF2CC"
PREENCHIMENTO_STATUS = "EDEDED"
PREENCHIMENTO_KPI = "DDEBF7"
PREENCHIMENTO_ALERTA = "FCE4D6"
PREENCHIMENTO_VERDE = "C6EFCE"
PREENCHIMENTO_LARANJA = "FCE4D6"
PREENCHIMENTO_VERMELHO = "FFC7CE"
PREENCHIMENTO_AMARELO = "FFF2CC"

STATUS_CANDIDATO = ["PENDENTE", "APROVADO", "AJUSTADO", "REJEITADO"]
STATUS_DOCUMENTO = ["PENDENTE_REVISAO", "VALIDADO_COMPLETO"]


def _linhas_como_dicionario(planilha):
    item = planilha.iter_rows(values_only=True)
    try:
        cabecalhos = list(next(item))
    except StopIteration:
        return
    for linha in item:
        yield {h: linha[i] if i < len(linha) else None for i, h in enumerate(cabecalhos)}


def _estilizar_cabecalho(planilha):
    for celula in planilha[1]:
        celula.font = Font(bold=True, color=FONTE_CABECALHO)
        celula.fill = PatternFill("solid", fgColor=PREENCHIMENTO_CABECALHO)
        celula.alignment = Alignment(vertical="top", wrap_text=True)
    planilha.freeze_panes = "A2"
    planilha.auto_filter.ref = planilha.dimensions


def _adicionar_lista_validacao(planilha, intervalo_celulas: str, formula: str):
    validacao_dados = DataValidation(type="list", formula1=formula, allow_blank=False)
    validacao_dados.error = "Escolha um valor da lista."
    validacao_dados.errorTitle = "Valor inválido"
    validacao_dados.prompt = "Selecione uma das opções disponíveis."
    validacao_dados.promptTitle = "Preenchimento padronizado"
    validacao_dados.showInputMessage = True
    validacao_dados.showErrorMessage = True
    planilha.add_data_validation(validacao_dados)
    validacao_dados.add(intervalo_celulas)


def _preencher_aba(planilha, linhas):
    if not linhas:
        return
    cabecalhos = list(linhas[0].keys())
    planilha.append(cabecalhos)
    for r in linhas:
        planilha.append([r.get(h, "") for h in cabecalhos])
    _estilizar_cabecalho(planilha)
    for indice, h in enumerate(cabecalhos, start=1):
        largura = 18
        if h in {"decisao", "candidate_text", "teacher_reason", "observacao", "observacao_documento", "texto_ajustado"}:
            largura = 55
        planilha.column_dimensions[get_column_letter(indice)].width = largura
    for linha in planilha.iter_rows():
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=False)


def _chave_decisao(d: dict) -> str:
    return str(d.get("decision_sha256") or f"source_row:{d.get('source_row')}")


def _construir_grupos(decisoes: list[dict], candidatos: list[dict]):
    decisoes_por_hash: dict[str, list[dict]] = defaultdict(list)
    for d in decisoes:
        decisoes_por_hash[_chave_decisao(d)].append(d)

    hash_por_linha_origem = {}
    for h, grupo_decisoes in decisoes_por_hash.items():
        for d in grupo_decisoes:
            hash_por_linha_origem[str(d.get("source_row"))] = h

    candidatos_por_hash: dict[str, list[dict]] = defaultdict(list)
    for c in candidatos:
        h = str(c.get("decision_sha256") or hash_por_linha_origem.get(str(c.get("source_row")), f"source_row:{c.get('source_row')}"))
        candidatos_por_hash[h].append(c)
    return decisoes_por_hash, candidatos_por_hash


def _anotacoes_anteriores(linhas_anteriores: list[dict] | None) -> dict[str, dict]:
    saida: dict[str, dict] = {}
    for r in linhas_anteriores or []:
        h = str(r.get("_decision_sha256") or r.get("decision_sha256") or "").strip()
        if h:
            saida[h] = r
    return saida


def _construir_linhas_revisao(
    decisoes: list[dict],
    candidatos: list[dict],
    quantidade_recortes: int,
    linhas_anteriores: list[dict] | None = None,
) -> tuple[list[dict], list[dict]]:
    """Cria a interface humana compacta e uma tabela técnica 1:1 por decisão.

    A interface mostra apenas contexto necessário + recortes/categorias editáveis + dropdowns.
    Toda a proveniência fica em dados_tecnicos, escondida dos revisores.
    """
    decisoes_por_hash, candidatos_por_hash = _construir_grupos(decisoes, candidatos)
    anotacoes_previas = _anotacoes_anteriores(linhas_anteriores)

    linhas_revisao: list[dict] = []
    linhas_tecnicas: list[dict] = []

    for sequencia, (h, grupo_decisoes) in enumerate(decisoes_por_hash.items(), start=1):
        representante = grupo_decisoes[0]
        id_decisao = f"D{sequencia:04d}"
        linhas_origem = [str(d.get("source_row")) for d in grupo_decisoes if d.get("source_row") not in (None, "")]
        processos_encontrados = sorted({str(d.get("processo") or d.get("Processo") or "") for d in grupo_decisoes if str(d.get("processo") or d.get("Processo") or "").strip()})
        expedientes_encontrados = sorted({str(d.get("expediente") or "") for d in grupo_decisoes if str(d.get("expediente") or "").strip()})
        prazos_encontrados = sorted({str(d.get("prazo") or "") for d in grupo_decisoes if str(d.get("prazo") or "").strip()})
        links_encontrados = sorted({str(d.get("link") or "") for d in grupo_decisoes if str(d.get("link") or "").strip()})

        candidatos_decisao = candidatos_por_hash.get(h, [])
        por_posicao: dict[int, dict] = {}
        for c in candidatos_decisao:
            try:
                posicao = int(c.get("candidate_rank") or 0)
            except Exception:
                posicao = 0
            if posicao > 0 and posicao not in por_posicao:
                por_posicao[posicao] = c

        anotacao_antiga = anotacoes_previas.get(h, {})
        revisao = {
            "_decision_sha256": h,
            "ID": id_decisao,
            "Processo": " | ".join(processos_encontrados),
            "Classe": representante.get("classe", ""),
            "Órgão julgador": representante.get("orgao_julgador", ""),
            "Expediente": " | ".join(expedientes_encontrados),
            "Prazo": " | ".join(prazos_encontrados),
            "Assunto": representante.get("assunto", ""),
        }

        # Candidatos: texto/categoria já são os campos finais editáveis.
        for i in range(1, quantidade_recortes + 1):
            c = por_posicao.get(i, {})
            texto_original = str(c.get("candidate_text") or "")
            categoria_original = str(c.get("teacher_category") or "")

            # Migração transparente da v4.2: preserva correções já feitas.
            status_antigo = str(anotacao_antiga.get(f"candidate_{i}_status") or anotacao_antiga.get(f"validacao_{i}") or "").strip().upper()
            texto_ajustado_antigo = str(anotacao_antiga.get(f"candidate_{i}_texto_ajustado") or "").strip()
            categoria_ajustada_antiga = str(anotacao_antiga.get(f"candidate_{i}_categoria_ajustada") or "").strip()
            texto_direto_antigo = anotacao_antiga.get(f"recorte_{i}")
            categoria_direta_antiga = anotacao_antiga.get(f"categoria_{i}")

            texto_final = str(texto_direto_antigo if texto_direto_antigo not in (None, "") else (texto_ajustado_antigo if status_antigo == "AJUSTADO" and texto_ajustado_antigo else texto_original))
            categoria_final = str(categoria_direta_antiga if categoria_direta_antiga not in (None, "") else (categoria_ajustada_antiga if status_antigo == "AJUSTADO" and categoria_ajustada_antiga else categoria_original))

            if texto_original:
                situacao = status_antigo if status_antigo in STATUS_CANDIDATO else "PENDENTE"
            else:
                situacao = ""

            revisao[f"Recorte {i}"] = texto_final
            revisao[f"Categoria {i}"] = categoria_final
            revisao[f"Validação {i}"] = situacao

        # Apenas dois espaços extras na tela principal; casos raros adicionais usam adicoes_gold.
        for i in range(1, 3):
            revisao[f"Recorte ausente {i}"] = str(anotacao_antiga.get(f"adicao_{i}_text") or anotacao_antiga.get(f"Recorte ausente {i}") or "")
            revisao[f"Categoria ausente {i}"] = str(anotacao_antiga.get(f"adicao_{i}_category") or anotacao_antiga.get(f"Categoria ausente {i}") or "")

        status_documento_antigo = str(anotacao_antiga.get("status_documento") or anotacao_antiga.get("Status da decisão") or "").strip().upper()
        revisao.update({
            "Status da decisão": status_documento_antigo if status_documento_antigo in STATUS_DOCUMENTO else "PENDENTE_REVISAO",
            "Revisor": str(anotacao_antiga.get("revisor") or anotacao_antiga.get("Revisor") or ""),
            "Observação": str(anotacao_antiga.get("observacao_documento") or anotacao_antiga.get("Observação") or ""),
        })
        linhas_revisao.append(revisao)

        dados_tecnicos = {
            "_decision_sha256": h,
            "ID": id_decisao,
            "source_rows": ", ".join(linhas_origem),
            "ocorrencias_na_base": len(grupo_decisoes),
            "Processo": " | ".join(processos_encontrados),
            "classe": representante.get("classe", ""),
            "origem": representante.get("origem", ""),
            "orgao_julgador": representante.get("orgao_julgador", ""),
            "data": representante.get("data", ""),
            "expedientes": " | ".join(expedientes_encontrados),
            "prazos_origem": " | ".join(prazos_encontrados),
            "assunto": representante.get("assunto", ""),
            "secao_subsecao": representante.get("secao_subsecao", ""),
            "nucleo": representante.get("nucleo", ""),
            "rule_id": representante.get("rule_id", ""),
            "rule_label": representante.get("rule_label", ""),
            "route_key": representante.get("route_key", ""),
            "specialist": representante.get("specialist", ""),
            "teacher_status": representante.get("teacher_status", ""),
            "teacher_attempts": representante.get("teacher_attempts", ""),
            "teacher_response_mode": representante.get("teacher_response_mode", ""),
            "cache_reutilizado": any(bool(d.get("cache_hit")) for d in grupo_decisoes) or len(grupo_decisoes) > 1,
            "links_origem": " | ".join(links_encontrados),
            "decisao": representante.get("decisao", ""),
        }
        for i in range(1, quantidade_recortes + 1):
            c = por_posicao.get(i, {})
            dados_tecnicos[f"teacher_text_{i}"] = c.get("candidate_text", "")
            dados_tecnicos[f"teacher_category_{i}"] = c.get("teacher_category", "")
            dados_tecnicos[f"teacher_priority_{i}"] = c.get("teacher_priority", "")
            dados_tecnicos[f"teacher_reason_{i}"] = c.get("teacher_reason", "")
        linhas_tecnicas.append(dados_tecnicos)

    return linhas_revisao, linhas_tecnicas


def _formatar_revisao(planilha, quantidade_recortes: int, quantidade_linhas: int):
    _estilizar_cabecalho(planilha)
    planilha.sheet_view.showGridLines = False
    planilha.freeze_panes = "C2"

    mapa_cabecalhos = {c.value: c.column for c in planilha[1]}

    # Coluna de hash existe para ligação técnica, mas fica invisível para o procurador.
    if "_decision_sha256" in mapa_cabecalhos:
        planilha.column_dimensions[get_column_letter(mapa_cabecalhos["_decision_sha256"])].hidden = True

    larguras = {
        "ID": 9,
        "Processo": 25,
        "Classe": 16,
        "Órgão julgador": 24,
        "Expediente": 28,
        "Prazo": 12,
        "Assunto": 28,
        "Status da decisão": 20,
        "Revisor": 20,
        "Observação": 32,
    }
    for nome, largura in larguras.items():
        if nome in mapa_cabecalhos:
            planilha.column_dimensions[get_column_letter(mapa_cabecalhos[nome])].width = largura

    for i in range(1, quantidade_recortes + 1):
        for nome, largura in {
            f"Recorte {i}": 48,
            f"Categoria {i}": 24,
            f"Validação {i}": 15,
        }.items():
            if nome in mapa_cabecalhos:
                planilha.column_dimensions[get_column_letter(mapa_cabecalhos[nome])].width = largura

    for i in range(1, 3):
        for nome, largura in {
            f"Recorte ausente {i}": 42,
            f"Categoria ausente {i}": 24,
        }.items():
            if nome in mapa_cabecalhos:
                planilha.column_dimensions[get_column_letter(mapa_cabecalhos[nome])].width = largura

    # Altura fixa: evita que textos longos criem linhas gigantes.
    for r in range(2, quantidade_linhas + 2):
        planilha.row_dimensions[r].height = 54

    # Sem wrap por padrão; campos de texto de revisão usam wrap dentro da altura fixa.
    for linha in planilha.iter_rows():
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=False)

    colunas_quebra_texto = ["Assunto", "Observação"]
    for i in range(1, quantidade_recortes + 1):
        colunas_quebra_texto.append(f"Recorte {i}")
    for i in range(1, 3):
        colunas_quebra_texto.append(f"Recorte ausente {i}")
    for nome in colunas_quebra_texto:
        if nome in mapa_cabecalhos:
            coluna = get_column_letter(mapa_cabecalhos[nome])
            for c in planilha[coluna][1:]:
                c.alignment = Alignment(vertical="top", wrap_text=True)

    # Cores por blocos para leitura rápida.
    colunas_contexto = ["ID", "Processo", "Classe", "Órgão julgador", "Expediente", "Prazo", "Assunto"]
    for nome in colunas_contexto:
        if nome in mapa_cabecalhos:
            coluna = get_column_letter(mapa_cabecalhos[nome])
            for c in planilha[coluna][1:]:
                c.fill = PatternFill("solid", fgColor=PREENCHIMENTO_CONTEXTO)

    for i in range(1, quantidade_recortes + 1):
        for nome in (f"Recorte {i}", f"Categoria {i}", f"Validação {i}"):
            if nome in mapa_cabecalhos:
                coluna = get_column_letter(mapa_cabecalhos[nome])
                for c in planilha[coluna][1:]:
                    c.fill = PatternFill("solid", fgColor=PREENCHIMENTO_CANDIDATO)

    for i in range(1, 3):
        for nome in (f"Recorte ausente {i}", f"Categoria ausente {i}"):
            if nome in mapa_cabecalhos:
                coluna = get_column_letter(mapa_cabecalhos[nome])
                for c in planilha[coluna][1:]:
                    c.fill = PatternFill("solid", fgColor=PREENCHIMENTO_HUMANO)

    for nome in ("Status da decisão", "Revisor", "Observação"):
        if nome in mapa_cabecalhos:
            coluna = get_column_letter(mapa_cabecalhos[nome])
            for c in planilha[coluna][1:]:
                c.fill = PatternFill("solid", fgColor=PREENCHIMENTO_STATUS)

    if quantidade_linhas:
        fim = quantidade_linhas + 1
        for i in range(1, quantidade_recortes + 1):
            categoria = get_column_letter(mapa_cabecalhos[f"Categoria {i}"])
            situacao = get_column_letter(mapa_cabecalhos[f"Validação {i}"])
            _adicionar_lista_validacao(planilha, f"{categoria}2:{categoria}{fim}", f"=listas!$B$2:$B${len(CATEGORIAS)+1}")
            # Status só é obrigatório para linhas que têm candidato; linhas vazias permanecem vazias.
            validacao_dados = DataValidation(type="list", formula1="=listas!$A$2:$A$5", allow_blank=True)
            validacao_dados.error = "Escolha PENDENTE, APROVADO, AJUSTADO ou REJEITADO."
            validacao_dados.showErrorMessage = True
            planilha.add_data_validation(validacao_dados)
            validacao_dados.add(f"{situacao}2:{situacao}{fim}")

            # Cores do dropdown de validação.
            planilha.conditional_formatting.add(
                f"{situacao}2:{situacao}{fim}",
                FormulaRule(formula=[f'${situacao}2="APROVADO"'], fill=PatternFill("solid", fgColor=PREENCHIMENTO_VERDE)),
            )
            planilha.conditional_formatting.add(
                f"{situacao}2:{situacao}{fim}",
                FormulaRule(formula=[f'${situacao}2="AJUSTADO"'], fill=PatternFill("solid", fgColor=PREENCHIMENTO_LARANJA)),
            )
            planilha.conditional_formatting.add(
                f"{situacao}2:{situacao}{fim}",
                FormulaRule(formula=[f'${situacao}2="REJEITADO"'], fill=PatternFill("solid", fgColor=PREENCHIMENTO_VERMELHO)),
            )
            planilha.conditional_formatting.add(
                f"{situacao}2:{situacao}{fim}",
                FormulaRule(formula=[f'${situacao}2="PENDENTE"'], fill=PatternFill("solid", fgColor=PREENCHIMENTO_AMARELO)),
            )

        for i in range(1, 3):
            categoria = get_column_letter(mapa_cabecalhos[f"Categoria ausente {i}"])
            validacao_dados = DataValidation(type="list", formula1=f"=listas!$B$2:$B${len(CATEGORIAS)+1}", allow_blank=True)
            validacao_dados.showErrorMessage = True
            planilha.add_data_validation(validacao_dados)
            validacao_dados.add(f"{categoria}2:{categoria}{fim}")

        status_documento = get_column_letter(mapa_cabecalhos["Status da decisão"])
        _adicionar_lista_validacao(planilha, f"{status_documento}2:{status_documento}{fim}", "=listas!$C$2:$C$3")
        planilha.conditional_formatting.add(
            f"{status_documento}2:{status_documento}{fim}",
            FormulaRule(formula=[f'${status_documento}2="VALIDADO_COMPLETO"'], fill=PatternFill("solid", fgColor=PREENCHIMENTO_VERDE)),
        )

        # Alerta visual: não marcar documento completo enquanto algum candidato ainda estiver PENDENTE.
        verificacoes_pendentes = []
        for i in range(1, quantidade_recortes + 1):
            coluna_texto = get_column_letter(mapa_cabecalhos[f"Recorte {i}"])
            coluna_validacao = get_column_letter(mapa_cabecalhos[f"Validação {i}"])
            verificacoes_pendentes.append(f'AND(${coluna_texto}2<>"",${coluna_validacao}2="PENDENTE")')
        if verificacoes_pendentes:
            formula = f'=AND(${status_documento}2="VALIDADO_COMPLETO",OR({",".join(verificacoes_pendentes)}))'
            planilha.conditional_formatting.add(
                f"A2:{get_column_letter(planilha.max_column)}{fim}",
                FormulaRule(formula=[formula], fill=PatternFill("solid", fgColor=PREENCHIMENTO_VERMELHO)),
            )


def _construir_aba_leitura(planilha, linhas_tecnicas: list[dict]):
    linhas = []
    for t in linhas_tecnicas:
        linhas.append({
            "ID": t.get("ID", ""),
            "Processo": t.get("Processo", ""),
            "Classe": t.get("classe", ""),
            "Assunto": t.get("assunto", ""),
            "Decisão completa": t.get("decisao", ""),
        })
    _preencher_aba(planilha, linhas)
    planilha.sheet_view.showGridLines = False
    planilha.freeze_panes = "A2"
    larguras = {"A": 10, "B": 26, "C": 18, "D": 28, "E": 100}
    for coluna, largura in larguras.items():
        planilha.column_dimensions[coluna].width = largura
    for r in range(2, planilha.max_row + 1):
        planilha.row_dimensions[r].height = 110
        planilha.cell(r, 5).alignment = Alignment(vertical="top", wrap_text=True)


def _construir_aba_metricas(planilha, aba_revisao, aba_tecnica, quantidade_recortes: int, linhas_revisao: int):
    planilha.sheet_view.showGridLines = False
    planilha["A1"] = "DIDE1 — Qualidade após revisão humana"
    planilha["A1"].font = Font(bold=True, size=14)
    planilha.merge_cells("A1:D1")
    planilha["A2"] = "As métricas só são confiáveis depois que o procurador revisa a decisão inteira e marca VALIDADO_COMPLETO."
    planilha.merge_cells("A2:D2")
    planilha["A2"].fill = PatternFill("solid", fgColor=PREENCHIMENTO_ALERTA)
    planilha["A2"].alignment = Alignment(wrap_text=True)

    h = {c.value: get_column_letter(c.column) for c in aba_revisao[1]}
    fim = max(2, linhas_revisao + 1)
    aba_atual = "'revisao'"

    def somar_status(valor: str) -> str:
        partes = []
        for i in range(1, quantidade_recortes + 1):
            c = h[f"Validação {i}"]
            partes.append(f'COUNTIF({aba_atual}!${c}$2:${c}${fim},"{valor}")')
        return "=" + "+".join(partes)

    candidato_preenchido = []
    for i in range(1, quantidade_recortes + 1):
        c = h[f"Recorte {i}"]
        candidato_preenchido.append(f'COUNTIF({aba_atual}!${c}$2:${c}${fim},"<>")')

    adicao_preenchida = []
    for i in range(1, 3):
        c = h[f"Recorte ausente {i}"]
        adicao_preenchida.append(f'COUNTIF({aba_atual}!${c}$2:${c}${fim},"<>")')

    linhas = [
        ["MÉTRICA", "VALOR", "INTERPRETAÇÃO"],
        ["Decisões únicas", f'=COUNTA({aba_atual}!${h["ID"]}$2:${h["ID"]}${fim})', "Uma linha por decisão única."],
        ["Decisões VALIDADO_COMPLETO", f'=COUNTIF({aba_atual}!${h["Status da decisão"]}$2:${h["Status da decisão"]}${fim},"VALIDADO_COMPLETO")', "Entram no GOLD somente após revisão completa."],
        ["Candidatos propostos", "=" + "+".join(candidato_preenchido), "Recortes sugeridos pelo teacher."],
        ["APROVADOS", somar_status("APROVADO"), "Corretos sem alteração humana."],
        ["AJUSTADOS", somar_status("AJUSTADO"), "Úteis, mas o procurador editou texto e/ou categoria."],
        ["REJEITADOS", somar_status("REJEITADO"), "Falsos positivos."],
        ["PENDENTES", somar_status("PENDENTE"), "Ainda não revisados."],
        ["Adições humanas", "=" + "+".join(adicao_preenchida), "Recortes relevantes que o modelo deixou escapar."],
    ]
    for linha in linhas:
        planilha.append(linha)
    for c in planilha[3]:
        c.fill = PatternFill("solid", fgColor=PREENCHIMENTO_CABECALHO)
        c.font = Font(bold=True, color=FONTE_CABECALHO)

    linhas_metricas = {planilha.cell(r, 1).value: r for r in range(1, planilha.max_row + 1)}
    aprovados = f"B{linhas_metricas['APROVADOS']}"
    ajustados = f"B{linhas_metricas['AJUSTADOS']}"
    rejeitados = f"B{linhas_metricas['REJEITADOS']}"
    adicoes = f"B{linhas_metricas['Adições humanas']}"
    validados = f"B{linhas_metricas['Decisões VALIDADO_COMPLETO']}"

    planilha.append([])
    planilha.append(["INDICADOR", "VALOR", "DEFINIÇÃO"])
    for c in planilha[planilha.max_row]:
        c.font = Font(bold=True, color="000000")
        c.fill = PatternFill("solid", fgColor=PREENCHIMENTO_KPI)

    linha_precisao = planilha.max_row + 1
    qualidade = [
        ["Precisão estrita", f'=IFERROR({aprovados}/({aprovados}+{ajustados}+{rejeitados}),"")', "Aprovados / candidatos efetivamente revisados."],
        ["Taxa útil", f'=IFERROR(({aprovados}+{ajustados})/({aprovados}+{ajustados}+{rejeitados}),"")', "Aprovados + ajustados / revisados."],
        ["Recall estimado", f'=IFERROR(({aprovados}+{ajustados})/({aprovados}+{ajustados}+{adicoes}),"")', "Recortes úteis encontrados / encontrados + adições humanas."],
    ]
    for linha in qualidade:
        planilha.append(linha)
    linha_recall = linha_precisao + 2
    planilha.append(["F1 estimado", f'=IFERROR(2*B{linha_precisao}*B{linha_recall}/(B{linha_precisao}+B{linha_recall}),"")', "Média harmônica entre precisão estrita e recall estimado."])

    # Acurácia de categoria: APROVADO é correto; AJUSTADO conta como correto se categoria final = teacher original.
    # Para manter a fórmula simples, o indicador é calculado na geração do GOLD e registrado no manifest do GOLD.
    planilha.append(["Cobertura completa", f'=IFERROR(SUMPRODUCT(--({aba_atual}!${h["Status da decisão"]}$2:${h["Status da decisão"]}${fim}="VALIDADO_COMPLETO"),--({aba_atual}!${h["Recorte ausente 1"]}$2:${h["Recorte ausente 1"]}${fim}=""),--({aba_atual}!${h["Recorte ausente 2"]}$2:${h["Recorte ausente 2"]}${fim}=""))/{validados},"")', "Documentos validados sem recortes adicionais humanos."])

    inicio = linha_precisao
    for r in range(inicio, planilha.max_row + 1):
        planilha.cell(r, 2).number_format = "0.0%"

    planilha.column_dimensions["A"].width = 32
    planilha.column_dimensions["B"].width = 18
    planilha.column_dimensions["C"].width = 80
    for linha in planilha.iter_rows():
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    planilha.freeze_panes = "A4"


def _construir_instrucoes(planilha):
    linhas = [
        ["PASSO", "O QUE FAZER"],
        ["1", "Abra 'leitura_decisoes' e leia a decisão integral pelo ID/Processo."],
        ["2", "Na aba 'revisao', confira Recorte 1/2/3. Em Validação escolha APROVADO, AJUSTADO ou REJEITADO."],
        ["3", "Se escolher AJUSTADO, edite diretamente o texto em 'Recorte N' e/ou a 'Categoria N'. O original do modelo fica preservado na aba técnica oculta."],
        ["4", "Se o modelo deixou um trecho relevante de fora, preencha 'Recorte ausente' e escolha a categoria no menu. Para mais de 2 ausências, use 'adicoes_gold'."],
        ["5", "Somente depois de ler a decisão inteira e revisar todos os recortes marque 'Status da decisão' = VALIDADO_COMPLETO."],
        ["6", "Não altere as abas técnicas ocultas. Elas guardam proveniência, texto integral e saída original para treinamento/auditoria."],
        ["GOLD", "O comando build-gold usa apenas decisões VALIDADO_COMPLETO e bloqueia documentos com candidatos ainda PENDENTE."],
    ]
    for r in linhas:
        planilha.append(r)
    _estilizar_cabecalho(planilha)
    planilha.column_dimensions["A"].width = 18
    planilha.column_dimensions["B"].width = 120
    for linha in planilha.iter_rows():
        for c in linha:
            c.alignment = Alignment(vertical="top", wrap_text=True)


def _construir_listas(planilha):
    planilha.append(["validacao", "categoria", "status_documento"])
    tamanho_maximo = max(len(STATUS_CANDIDATO), len(CATEGORIAS), len(STATUS_DOCUMENTO))
    for i in range(tamanho_maximo):
        planilha.append([
            STATUS_CANDIDATO[i] if i < len(STATUS_CANDIDATO) else "",
            CATEGORIAS[i] if i < len(CATEGORIAS) else "",
            STATUS_DOCUMENTO[i] if i < len(STATUS_DOCUMENTO) else "",
        ])


def escrever_planilha_revisao(caminho, decisoes, candidatos, linhas_anteriores: list[dict] | None = None):
    quantidade_recortes = max(3, max([int(c.get("candidate_rank") or 0) for c in candidatos] or [0]))
    linhas_revisao, linhas_tecnicas = _construir_linhas_revisao(decisoes, candidatos, quantidade_recortes, linhas_anteriores=linhas_anteriores)

    pasta_trabalho = Workbook()
    revisao = pasta_trabalho.active
    revisao.title = "revisao"
    if linhas_revisao:
        cabecalhos = list(linhas_revisao[0].keys())
        revisao.append(cabecalhos)
        for r in linhas_revisao:
            revisao.append([r.get(h, "") for h in cabecalhos])
        _formatar_revisao(revisao, quantidade_recortes, len(linhas_revisao))

    leitura = pasta_trabalho.create_sheet("leitura_decisoes")
    _construir_aba_leitura(leitura, linhas_tecnicas)

    metricas = pasta_trabalho.create_sheet("metricas")
    tecnico = pasta_trabalho.create_sheet("dados_tecnicos")
    _preencher_aba(tecnico, linhas_tecnicas)
    tecnico.sheet_state = "hidden"

    aba_candidatos = pasta_trabalho.create_sheet("candidatos_tecnicos")
    _preencher_aba(aba_candidatos, candidatos)
    aba_candidatos.sheet_state = "hidden"

    instrucoes = pasta_trabalho.create_sheet("instrucoes")
    _construir_instrucoes(instrucoes)

    aba_adicoes = pasta_trabalho.create_sheet("adicoes_gold")
    aba_adicoes.append(["ID", "decision_sha256", "Processo", "texto_gold", "categoria_gold", "observacao", "revisor", "revisado_em"])
    _estilizar_cabecalho(aba_adicoes)
    aba_adicoes.column_dimensions["D"].width = 80
    aba_adicoes.column_dimensions["E"].width = 28
    validacao_dados = DataValidation(type="list", formula1=f"=listas!$B$2:$B${len(CATEGORIAS)+1}", allow_blank=True)
    aba_adicoes.add_data_validation(validacao_dados)
    validacao_dados.add("E2:E5000")

    listas_validacao = pasta_trabalho.create_sheet("listas")
    _construir_listas(listas_validacao)
    listas_validacao.sheet_state = "hidden"

    _construir_aba_metricas(metricas, revisao, tecnico, quantidade_recortes, len(linhas_revisao))

    pasta_trabalho.active = pasta_trabalho.sheetnames.index("revisao")
    pasta_trabalho.save(Path(caminho))


def _ler_linhas_revisao_anterior(pasta_trabalho) -> list[dict]:
    if "revisao" in pasta_trabalho.sheetnames:
        return list(_linhas_como_dicionario(pasta_trabalho["revisao"]))
    if "revisao_consolidada" in pasta_trabalho.sheetnames:
        return list(_linhas_como_dicionario(pasta_trabalho["revisao_consolidada"]))
    return []


def reconstruir_revisao(planilha_revisao_xlsx: str, arquivo_saida_xlsx: str):
    """Recria somente a planilha de revisão, sem chamar o Qwen.

    Aceita workbook v4.2/v4.3. Preserva anotações humanas já existentes quando possível.
    """
    pasta_trabalho = load_workbook(planilha_revisao_xlsx, data_only=False)
    anotacoes_previas = _ler_linhas_revisao_anterior(pasta_trabalho)

    # Preferir abas técnicas originais da v4.2; na v4.3 usar as equivalentes ocultas.
    if "decisoes" in pasta_trabalho.sheetnames:
        decisoes = list(_linhas_como_dicionario(pasta_trabalho["decisoes"]))
    elif "dados_tecnicos" in pasta_trabalho.sheetnames:
        # Reconstrói formato mínimo de decisões a partir da tabela 1:1 técnica.
        decisoes = []
        for t in _linhas_como_dicionario(pasta_trabalho["dados_tecnicos"]):
            decisoes.append({
                "source_row": str(t.get("source_rows") or "").split(",")[0].strip() or None,
                "processo": t.get("Processo", ""),
                "classe": t.get("classe", ""),
                "origem": t.get("origem", ""),
                "orgao_julgador": t.get("orgao_julgador", ""),
                "data": t.get("data", ""),
                "expediente": t.get("expedientes", ""),
                "prazo": t.get("prazos_origem", ""),
                "assunto": t.get("assunto", ""),
                "secao_subsecao": t.get("secao_subsecao", ""),
                "nucleo": t.get("nucleo", ""),
                "link": t.get("links_origem", ""),
                "decisao": t.get("decisao", ""),
                "rule_id": t.get("rule_id", ""),
                "rule_label": t.get("rule_label", ""),
                "route_key": t.get("route_key", ""),
                "specialist": t.get("specialist", ""),
                "teacher_status": t.get("teacher_status", ""),
                "teacher_attempts": t.get("teacher_attempts", ""),
                "teacher_response_mode": t.get("teacher_response_mode", ""),
                "decision_sha256": t.get("_decision_sha256", ""),
                "cache_hit": bool(t.get("cache_reutilizado")),
            })
    else:
        raise ValueError("Workbook não contém aba 'decisoes' nem 'dados_tecnicos'.")

    if "candidatos" in pasta_trabalho.sheetnames:
        candidatos = list(_linhas_como_dicionario(pasta_trabalho["candidatos"]))
    elif "candidatos_tecnicos" in pasta_trabalho.sheetnames:
        candidatos = list(_linhas_como_dicionario(pasta_trabalho["candidatos_tecnicos"]))
    else:
        raise ValueError("Workbook não contém aba 'candidatos' nem 'candidatos_tecnicos'.")

    # Preserva adições extras existentes e o antigo terceiro slot da v4.2.
    adicoes_preservadas = []
    if "adicoes_gold" in pasta_trabalho.sheetnames:
        for a in _linhas_como_dicionario(pasta_trabalho["adicoes_gold"]):
            texto = str(a.get("texto_gold") or "").strip()
            if texto:
                adicoes_preservadas.append(dict(a))
    for r in anotacoes_previas:
        texto = str(r.get("adicao_3_text") or "").strip()
        if texto:
            adicoes_preservadas.append({
                "decision_sha256": str(r.get("decision_sha256") or r.get("_decision_sha256") or ""),
                "Processo": r.get("Processo", ""),
                "texto_gold": texto,
                "categoria_gold": str(r.get("adicao_3_category") or ""),
                "observacao": "Migrado automaticamente do 3º recorte ausente da v4.2",
                "revisor": r.get("revisor", ""),
                "revisado_em": r.get("revisado_em", ""),
            })

    escrever_planilha_revisao(arquivo_saida_xlsx, decisoes, candidatos, linhas_anteriores=anotacoes_previas)

    if adicoes_preservadas:
        pasta_trabalho_saida = load_workbook(arquivo_saida_xlsx)
        aba_adicoes = pasta_trabalho_saida["adicoes_gold"]
        aba_tecnica = pasta_trabalho_saida["dados_tecnicos"]
        linhas_tecnicas = list(_linhas_como_dicionario(aba_tecnica))
        id_por_hash = {str(t.get("_decision_sha256") or ""): str(t.get("ID") or "") for t in linhas_tecnicas}
        for a in adicoes_preservadas:
            h = str(a.get("decision_sha256") or "").strip()
            aba_adicoes.append([
                id_por_hash.get(h, str(a.get("ID") or "")),
                h,
                a.get("Processo", ""),
                a.get("texto_gold", ""),
                a.get("categoria_gold", ""),
                a.get("observacao", ""),
                a.get("revisor", ""),
                a.get("revisado_em", ""),
            ])
        pasta_trabalho_saida.save(arquivo_saida_xlsx)

    return {
        "input": str(planilha_revisao_xlsx),
        "output": str(arquivo_saida_xlsx),
        "decisions": len(decisoes),
        "candidates": len(candidatos),
        "preserved_extra_additions": len(adicoes_preservadas),
        "qwen_rerun": False,
    }


def _mapa_tecnico(pasta_trabalho) -> dict[str, dict]:
    if "dados_tecnicos" not in pasta_trabalho.sheetnames:
        return {}
    saida = {}
    for t in _linhas_como_dicionario(pasta_trabalho["dados_tecnicos"]):
        h = str(t.get("_decision_sha256") or "")
        if h:
            saida[h] = t
    return saida


def _construir_gold_v43(pasta_trabalho, arquivo_saida_jsonl: str):
    revisao = pasta_trabalho["revisao"]
    linhas = list(_linhas_como_dicionario(revisao))
    tecnico = _mapa_tecnico(pasta_trabalho)
    cabecalhos = [c.value for c in revisao[1]]
    quantidade_recortes = 0
    while f"Recorte {quantidade_recortes + 1}" in cabecalhos:
        quantidade_recortes += 1

    adicoes_extras = defaultdict(list)
    if "adicoes_gold" in pasta_trabalho.sheetnames:
        for a in _linhas_como_dicionario(pasta_trabalho["adicoes_gold"]):
            texto = str(a.get("texto_gold") or "").strip()
            if not texto:
                continue
            h = str(a.get("decision_sha256") or "").strip()
            id_decisao_ref = str(a.get("ID") or "").strip()
            chave = h or id_decisao_ref
            adicoes_extras[chave].append({
                "text": texto,
                "category": str(a.get("categoria_gold") or "").strip(),
                "source": "human_addition",
            })

    validados = []
    problemas = []
    total_aprovados = total_ajustados = total_rejeitados = 0
    categorias_corretas = total_categorias = 0

    for linha_excel, r in enumerate(linhas, start=2):
        if str(r.get("Status da decisão") or "").strip().upper() != "VALIDADO_COMPLETO":
            continue

        h = str(r.get("_decision_sha256") or "").strip()
        t = tecnico.get(h, {})
        problemas_linha = []
        gold = []

        for i in range(1, quantidade_recortes + 1):
            texto = str(r.get(f"Recorte {i}") or "").strip()
            categoria = str(r.get(f"Categoria {i}") or "").strip()
            situacao = str(r.get(f"Validação {i}") or "").strip().upper()
            texto_modelo = str(t.get(f"teacher_text_{i}") or "").strip()
            categoria_modelo = str(t.get(f"teacher_category_{i}") or "").strip()

            # Slot sem candidato original não exige validação.
            if not texto_modelo and not texto:
                continue
            if situacao not in {"APROVADO", "AJUSTADO", "REJEITADO"}:
                problemas_linha.append(f"candidato {i} ainda não revisado")
                continue
            if situacao == "REJEITADO":
                total_rejeitados += 1
                continue
            if not texto:
                problemas_linha.append(f"candidato {i} sem texto final")
                continue
            if categoria not in CATEGORIAS:
                problemas_linha.append(f"candidato {i} com categoria inválida")
                continue

            if situacao == "APROVADO":
                total_aprovados += 1
            else:
                total_ajustados += 1
            total_categorias += 1
            if categoria == categoria_modelo:
                categorias_corretas += 1
            gold.append({"text": texto, "category": categoria, "source": "candidate", "validation": situacao})

        for i in range(1, 3):
            texto = str(r.get(f"Recorte ausente {i}") or "").strip()
            categoria = str(r.get(f"Categoria ausente {i}") or "").strip()
            if not texto:
                continue
            if categoria not in CATEGORIAS:
                problemas_linha.append(f"recorte ausente {i} sem categoria válida")
                continue
            gold.append({"text": texto, "category": categoria, "source": "human_addition"})

        id_decisao_ref = str(r.get("ID") or "")
        for chave in (h, id_decisao_ref):
            for a in adicoes_extras.get(chave, []):
                if a["category"] not in CATEGORIAS:
                    problemas_linha.append("adicao_gold com categoria inválida")
                else:
                    gold.append(a)

        if problemas_linha:
            problemas.append({"excel_row": linha_excel, "ID": id_decisao_ref, "Processo": r.get("Processo"), "issues": "; ".join(problemas_linha)})
            continue

        linhas_origem = [x.strip() for x in str(t.get("source_rows") or "").split(",") if x.strip()]
        objeto = {
            "source_row": int(linhas_origem[0]) if linhas_origem and linhas_origem[0].isdigit() else (linhas_origem[0] if linhas_origem else None),
            "source_rows": linhas_origem,
            "decision_sha256": h,
            "Processo": r.get("Processo"),
            "metadata": {
                "classe": t.get("classe", r.get("Classe", "")),
                "origem": t.get("origem", ""),
                "orgao_julgador": t.get("orgao_julgador", r.get("Órgão julgador", "")),
                "data": t.get("data", ""),
                "expediente": t.get("expedientes", r.get("Expediente", "")),
                "prazo": t.get("prazos_origem", r.get("Prazo", "")),
                "assunto": t.get("assunto", r.get("Assunto", "")),
                "secao_subsecao": t.get("secao_subsecao", ""),
                "nucleo": t.get("nucleo", ""),
                "route_key": t.get("route_key", ""),
                "specialist": t.get("specialist", ""),
            },
            "decision": t.get("decisao", ""),
            "gold": gold,
        }
        validados.append(objeto)

    saida = Path(arquivo_saida_jsonl)
    saida.parent.mkdir(parents=True, exist_ok=True)
    with saida.open("w", encoding="utf-8") as f:
        for objeto in validados:
            f.write(json.dumps(objeto, ensure_ascii=False, default=str) + "\n")

    caminho_problemas = saida.with_name("gold_build_issues.csv")
    with caminho_problemas.open("w", encoding="utf-8-sig", newline="") as f:
        gravador = csv.DictWriter(f, fieldnames=["excel_row", "ID", "Processo", "issues"])
        gravador.writeheader()
        gravador.writerows(problemas)

    metricas = {
        "validated_documents": len(validados),
        "blocked_documents": len(problemas),
        "approved": total_aprovados,
        "adjusted": total_ajustados,
        "rejected": total_rejeitados,
        "category_accuracy_on_useful": (categorias_corretas / total_categorias) if total_categorias else None,
        "output": str(saida),
        "issues": str(caminho_problemas),
        "source": "revisao",
    }
    return metricas


def _construir_gold_legado(pasta_trabalho, arquivo_saida_jsonl: str):
    # Compatibilidade com workbooks v4.0/v4.1/v4.2.
    if "revisao_consolidada" in pasta_trabalho.sheetnames:
        revisao = pasta_trabalho["revisao_consolidada"]
        linhas = list(_linhas_como_dicionario(revisao))
        cabecalhos = [c.value for c in revisao[1]]
        quantidade_recortes = 0
        while f"candidate_{quantidade_recortes+1}_text" in cabecalhos:
            quantidade_recortes += 1
        validados = []
        for r in linhas:
            if str(r.get("status_documento") or "").strip().upper() != "VALIDADO_COMPLETO":
                continue
            gold = []
            for i in range(1, quantidade_recortes + 1):
                situacao = str(r.get(f"candidate_{i}_status") or "").strip().upper()
                if situacao == "APROVADO":
                    texto = str(r.get(f"candidate_{i}_text") or "")
                    categoria = str(r.get(f"candidate_{i}_category") or "")
                elif situacao == "AJUSTADO":
                    texto = str(r.get(f"candidate_{i}_texto_ajustado") or r.get(f"candidate_{i}_text") or "")
                    categoria = str(r.get(f"candidate_{i}_categoria_ajustada") or r.get(f"candidate_{i}_category") or "")
                else:
                    continue
                if texto.strip():
                    gold.append({"text": texto, "category": categoria, "source": "candidate"})
            for i in range(1, 4):
                texto = str(r.get(f"adicao_{i}_text") or "").strip()
                if texto:
                    gold.append({"text": texto, "category": str(r.get(f"adicao_{i}_category") or "outro_acionavel"), "source": "human_addition"})
            linhas_origem = [x.strip() for x in str(r.get("source_rows") or "").split(",") if x.strip()]
            validados.append({
                "source_row": int(linhas_origem[0]) if linhas_origem and linhas_origem[0].isdigit() else None,
                "source_rows": linhas_origem,
                "decision_sha256": r.get("decision_sha256", ""),
                "Processo": r.get("Processo"),
                "metadata": {
                    "classe": r.get("classe", ""),
                    "origem": r.get("origem", ""),
                    "orgao_julgador": r.get("orgao_julgador", ""),
                    "data": r.get("data", ""),
                    "expediente": r.get("expedientes", ""),
                    "prazo": r.get("prazos_origem", ""),
                    "assunto": r.get("assunto", ""),
                    "secao_subsecao": r.get("secao_subsecao", ""),
                    "nucleo": r.get("nucleo", ""),
                    "route_key": r.get("route_key", ""),
                    "specialist": r.get("specialist", ""),
                },
                "decision": r.get("decisao", ""),
                "gold": gold,
            })
        saida = Path(arquivo_saida_jsonl)
        saida.parent.mkdir(parents=True, exist_ok=True)
        with saida.open("w", encoding="utf-8") as f:
            for objeto in validados:
                f.write(json.dumps(objeto, ensure_ascii=False, default=str) + "\n")
        return {"validated_documents": len(validados), "output": str(saida), "source": "revisao_consolidada"}

    decisoes = pasta_trabalho["decisoes"]
    aba_candidatos = pasta_trabalho["candidatos"]
    aba_adicoes = pasta_trabalho["adicoes_gold"]
    mapa_decisoes = {}
    for d in _linhas_como_dicionario(decisoes):
        if str(d.get("status_documento", "")).strip().upper() == "VALIDADO_COMPLETO":
            mapa_decisoes[int(d["source_row"])] = d
    gold = {k: [] for k in mapa_decisoes}
    for c in _linhas_como_dicionario(aba_candidatos):
        try:
            linha_origem_ref = int(c.get("source_row"))
        except Exception:
            continue
        if linha_origem_ref not in mapa_decisoes:
            continue
        situacao = str(c.get("status_validacao", "")).strip().upper()
        if situacao == "APROVADO":
            texto = str(c.get("candidate_text") or "")
            categoria = str(c.get("teacher_category") or "")
        elif situacao == "AJUSTADO":
            texto = str(c.get("texto_ajustado") or c.get("candidate_text") or "")
            categoria = str(c.get("categoria_ajustada") or c.get("teacher_category") or "")
        else:
            continue
        if texto.strip():
            gold[linha_origem_ref].append({"text": texto, "category": categoria, "source": "candidate"})
    for a in _linhas_como_dicionario(aba_adicoes):
        try:
            linha_origem_ref = int(a.get("source_row"))
        except Exception:
            continue
        if linha_origem_ref in mapa_decisoes and str(a.get("texto_gold") or "").strip():
            gold[linha_origem_ref].append({"text": str(a["texto_gold"]), "category": str(a.get("categoria_gold") or "outro_acionavel"), "source": "human_addition"})
    saida = Path(arquivo_saida_jsonl)
    saida.parent.mkdir(parents=True, exist_ok=True)
    with saida.open("w", encoding="utf-8") as f:
        for linha_origem_ref, d in mapa_decisoes.items():
            f.write(json.dumps({
                "source_row": linha_origem_ref,
                "Processo": d.get("processo") or d.get("Processo"),
                "metadata": {k: v for k, v in d.items() if k not in {"decisao", "status_documento"}},
                "decision": d.get("decisao", ""),
                "gold": gold.get(linha_origem_ref, []),
            }, ensure_ascii=False, default=str) + "\n")
    return {"validated_documents": len(mapa_decisoes), "output": str(saida), "source": "legacy"}


def construir_gold(planilha_revisao_xlsx: str, arquivo_saida_jsonl: str):
    pasta_trabalho = load_workbook(planilha_revisao_xlsx, data_only=True)
    if "revisao" in pasta_trabalho.sheetnames and "dados_tecnicos" in pasta_trabalho.sheetnames:
        return _construir_gold_v43(pasta_trabalho, arquivo_saida_jsonl)
    return _construir_gold_legado(pasta_trabalho, arquivo_saida_jsonl)
