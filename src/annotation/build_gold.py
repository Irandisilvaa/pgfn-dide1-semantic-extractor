from __future__ import annotations

import argparse
from collections import defaultdict
from dataclasses import asdict
import csv
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import random
import re
from typing import Any

from openpyxl import load_workbook

from src.annotation.workbook import ALLOWED_CATEGORIES, row_to_dict
from src.semantic.action_extractor import SYSTEM_PROMPT, CHUNK_PROMPT, FINAL_PROMPT
from src.semantic.deterministic_filters import low_value_reason
from src.semantic.segmenter import group_units, make_units


def norm(text: str) -> str:
    return re.sub(r"\s+", " ", str(text or "").strip()).casefold()


def phash(value: str, salt: str = "dide1-gold") -> str:
    return hashlib.sha256((salt + "|" + value).encode("utf-8")).hexdigest()[:16]


def parse_ids(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(x) for x in value]
    text = str(value or "").strip()
    if not text:
        return []
    try:
        obj = json.loads(text)
        if isinstance(obj, list):
            return [str(x) for x in obj]
    except Exception:
        pass
    return [x.strip() for x in text.split(",") if x.strip()]


def exact_unit_for_text(decision: str, text: str, *, max_unit_chars: int) -> tuple[str | None, str | None]:
    target = norm(text)
    matches = [u for u in make_units(decision, max_unit_chars=max_unit_chars) if norm(u.text) == target]
    if len(matches) == 1:
        return matches[0].id, matches[0].text
    if len(matches) > 1:
        return None, "AMBIGUOUS_UNIT_MATCH"
    return None, "NO_SINGLE_UNIT_MATCH"


def split_process_groups(records: list[dict], seed: int, train_ratio: float, val_ratio: float) -> dict[str, str]:
    groups: dict[str, list[dict]] = defaultdict(list)
    for r in records:
        key = str(r.get("Processo") or f"ROW-{r['source_row']}")
        groups[key].append(r)
    keys = sorted(groups)
    rng = random.Random(seed)
    rng.shuffle(keys)
    total_docs = len(records)
    target_train = total_docs * train_ratio
    target_val = total_docs * val_ratio
    counts = {"train": 0, "val": 0, "test": 0}
    assignment: dict[str, str] = {}
    for key in keys:
        n = len(groups[key])
        if counts["train"] < target_train:
            split = "train"
        elif counts["val"] < target_val:
            split = "val"
        else:
            split = "test"
        assignment[key] = split
        counts[split] += n
    return assignment


def main() -> None:
    p = argparse.ArgumentParser(description="Planilha humana validada -> DIDE1-GOLD + datasets SFT, com split por Processo.")
    p.add_argument("--review-xlsx", required=True)
    p.add_argument("--output-dir", default="runtime/private_gold")
    p.add_argument("--seed", type=int, default=20260916)
    p.add_argument("--train-ratio", type=float, default=0.80)
    p.add_argument("--val-ratio", type=float, default=0.10)
    p.add_argument("--max-unit-chars", type=int, default=1200)
    p.add_argument("--max-chunk-chars", type=int, default=5200)
    p.add_argument("--max-candidates-per-chunk", type=int, default=4)
    p.add_argument("--negative-chunk-ratio", type=float, default=1.0)
    args = p.parse_args()

    review = Path(args.review_xlsx)
    out = Path(args.output_dir)
    out.mkdir(parents=True, exist_ok=True)
    wb = load_workbook(review, read_only=True, data_only=False)

    ws_dec = wb["decisoes"]
    dec_headers = [str(c.value or "") for c in ws_dec[1]]
    decisions: dict[int, dict] = {}
    for values in ws_dec.iter_rows(min_row=2, values_only=True):
        row = row_to_dict(dec_headers, values)
        if not row.get("source_row"):
            continue
        decisions[int(row["source_row"])] = row

    ws_can = wb["candidatos"]
    can_headers = [str(c.value or "") for c in ws_can[1]]
    candidates: dict[int, list[dict]] = defaultdict(list)
    for values in ws_can.iter_rows(min_row=2, values_only=True):
        row = row_to_dict(can_headers, values)
        if row.get("source_row"):
            candidates[int(row["source_row"])].append(row)

    additions: dict[int, list[dict]] = defaultdict(list)
    if "adicoes_gold" in wb.sheetnames:
        ws_add = wb["adicoes_gold"]
        add_headers = [str(c.value or "") for c in ws_add[1]]
        for values in ws_add.iter_rows(min_row=2, values_only=True):
            row = row_to_dict(add_headers, values)
            if row.get("source_row") and str(row.get("gold_text_adicionado") or "").strip():
                additions[int(row["source_row"])].append(row)

    issues: list[dict[str, Any]] = []
    gold_records: list[dict[str, Any]] = []

    for source_row, dec in decisions.items():
        doc_status = str(dec.get("status_documento") or "").strip().upper()
        if doc_status not in {"VALIDADO_COMPLETO", "SEM_RECORTE"}:
            continue

        decision = str(dec.get("Decisão") or "")
        processo = str(dec.get("Processo") or "")
        doc_candidates = candidates.get(source_row, [])
        gold: list[dict] = []
        invalid = False

        if doc_status == "VALIDADO_COMPLETO":
            for cand in doc_candidates:
                rank = int(cand.get("candidate_rank") or 0)
                status = str(cand.get("status_validacao") or "").strip().upper()
                if rank == 0 and not str(cand.get("candidate_text") or "").strip():
                    if status not in {"", "SEM_RECORTE", "REJEITADO"}:
                        issues.append({"source_row": source_row, "issue": "INVALID_EMPTY_CANDIDATE_STATUS", "detail": status})
                        invalid = True
                    continue
                if status not in {"APROVADO", "AJUSTADO", "REJEITADO"}:
                    issues.append({"source_row": source_row, "issue": "CANDIDATE_NOT_REVIEWED", "detail": f"rank={rank}"})
                    invalid = True
                    continue
                if status == "REJEITADO":
                    continue
                if status == "APROVADO":
                    text = str(cand.get("candidate_text") or "").strip()
                    category = str(cand.get("teacher_category") or "").strip()
                else:
                    text = str(cand.get("texto_ajustado") or "").strip()
                    category = str(cand.get("categoria_ajustada") or "").strip()
                if not text or category not in ALLOWED_CATEGORIES:
                    issues.append({"source_row": source_row, "issue": "INVALID_GOLD_FIELDS", "detail": f"rank={rank}"})
                    invalid = True
                    continue
                uid, literal = exact_unit_for_text(decision, text, max_unit_chars=args.max_unit_chars)
                if uid is None:
                    issues.append({"source_row": source_row, "issue": literal, "detail": text[:160]})
                    invalid = True
                    continue
                gold.append({"text": literal, "category": category, "unit_id": uid, "origin": status.lower()})

            for add in additions.get(source_row, []):
                text = str(add.get("gold_text_adicionado") or "").strip()
                category = str(add.get("gold_category_adicionada") or "").strip()
                if category not in ALLOWED_CATEGORIES:
                    issues.append({"source_row": source_row, "issue": "INVALID_ADDITION_CATEGORY", "detail": category})
                    invalid = True
                    continue
                uid, literal = exact_unit_for_text(decision, text, max_unit_chars=args.max_unit_chars)
                if uid is None:
                    issues.append({"source_row": source_row, "issue": literal, "detail": text[:160]})
                    invalid = True
                    continue
                gold.append({"text": literal, "category": category, "unit_id": uid, "origin": "human_addition"})
        else:
            if additions.get(source_row):
                issues.append({"source_row": source_row, "issue": "SEM_RECORTE_WITH_ADDITIONS", "detail": ""})
                invalid = True

        if invalid:
            continue

        # dedup gold by unit id
        dedup: dict[str, dict] = {}
        for g in gold:
            dedup[g["unit_id"]] = g
        gold = list(dedup.values())

        gold_records.append({
            "source_row": source_row,
            "Processo": processo,
            "Classe judicial": dec.get("Classe judicial") or "",
            "Órgão julgador": dec.get("Órgão julgador") or "",
            "Polo Ativo": dec.get("Polo Ativo") or "",
            "Polo Passivo": dec.get("Polo Passivo") or "",
            "Matéria SAJ": dec.get("Matéria SAJ") or "",
            "Decisão": decision,
            "gold": gold,
            "teacher_candidates": doc_candidates,
            "doc_status": doc_status,
        })

    if not gold_records:
        raise SystemExit("Nenhum documento totalmente validado e sem inconsistências foi encontrado.")

    assignments = split_process_groups(gold_records, args.seed, args.train_ratio, args.val_ratio)
    for rec in gold_records:
        process_key = str(rec.get("Processo") or f"ROW-{rec['source_row']}")
        rec["process_hash"] = phash(process_key)
        rec["split"] = assignments[process_key]

    gold_path = out / "gold_records.jsonl"
    with gold_path.open("w", encoding="utf-8") as f:
        for rec in gold_records:
            clean = {k: v for k, v in rec.items() if k != "teacher_candidates"}
            f.write(json.dumps(clean, ensure_ascii=False) + "\n")

    # Build SFT examples: chunk selector + final adjudication where possible.
    rng = random.Random(args.seed)
    sft = {"train": [], "val": [], "test": []}
    for rec in gold_records:
        decision = rec["Decisão"]
        gold_by_id = {g["unit_id"]: g for g in rec["gold"]}
        units = make_units(decision, max_unit_chars=args.max_unit_chars)
        by_id = {u.id: u for u in units}
        for gid in gold_by_id:
            if gid not in by_id:
                issues.append({"source_row": rec["source_row"], "issue": "GOLD_UNIT_DISAPPEARED", "detail": gid})
        chunks = group_units(units, max_chunk_chars=args.max_chunk_chars)
        positives, negatives = [], []
        for chunk in chunks:
            selectable = [u for u in chunk if not low_value_reason(u.text)]
            ids = [u.id for u in selectable]
            filtered_gold = [gid for gid in gold_by_id if gid in {u.id for u in chunk} and gid not in ids]
            if filtered_gold:
                issues.append({"source_row": rec["source_row"], "issue": "GOLD_FILTERED_BY_PIPELINE", "detail": ",".join(filtered_gold)})
                continue
            rendered = "\n\n".join(f"[{u.id}][SECAO={u.section}] {u.text}" for u in selectable)
            selected = [
                {"id": gid, "categoria": gold_by_id[gid]["category"], "prioridade": 5}
                for gid in ids if gid in gold_by_id
            ]
            example = {
                "task": "chunk_selection",
                "process_hash": rec["process_hash"],
                "source_row": rec["source_row"],
                "messages": [
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": CHUNK_PROMPT.format(max_candidates=args.max_candidates_per_chunk, units=rendered)},
                    {"role": "assistant", "content": json.dumps({"recortes": selected}, ensure_ascii=False, separators=(",", ":"))},
                ],
            }
            (positives if selected else negatives).append(example)
        max_neg = max(1, int(round(len(positives) * args.negative_chunk_ratio))) if positives else 1
        if len(negatives) > max_neg:
            negatives = rng.sample(negatives, max_neg)
        sft[rec["split"]].extend(positives + negatives)

        # Final adjudication examples only when every original candidate has APROVADO/REJEITADO and no manual additions/adjustments.
        cands = rec.get("teacher_candidates", [])
        usable = []
        has_adjust = False
        for cand in cands:
            rank = int(cand.get("candidate_rank") or 0)
            if rank == 0:
                continue
            status = str(cand.get("status_validacao") or "").strip().upper()
            if status == "AJUSTADO":
                has_adjust = True
                break
            if status in {"APROVADO", "REJEITADO"}:
                usable.append(cand)
            else:
                has_adjust = True
                break
        human_added = any(g.get("origin") == "human_addition" for g in rec["gold"])
        if usable and not has_adjust and not human_added:
            rendered = []
            selected_cids = []
            for i, cand in enumerate(usable, start=1):
                cid = f"C{i:02d}"
                rendered.append(
                    f"[{cid}] SECAO={cand.get('section','')}; CATEGORIA={cand.get('teacher_category','')}; "
                    f"PRIORIDADE={cand.get('teacher_priority','')}; SCORE={cand.get('heuristic_score','')}\n{cand.get('candidate_text','')}"
                )
                if str(cand.get("status_validacao") or "").strip().upper() == "APROVADO":
                    selected_cids.append(cid)
            sft[rec["split"]].append({
                "task": "final_adjudication",
                "process_hash": rec["process_hash"],
                "source_row": rec["source_row"],
                "messages": [
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": FINAL_PROMPT.format(top_k=3, candidates="\n\n".join(rendered))},
                    {"role": "assistant", "content": json.dumps({"selecionados": selected_cids[:3]}, ensure_ascii=False, separators=(",", ":"))},
                ],
            })

    for split, rows in sft.items():
        path = out / f"sft_{split}.jsonl"
        with path.open("w", encoding="utf-8") as f:
            for row in rows:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")

    issues_path = out / "gold_build_issues.csv"
    with issues_path.open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["source_row", "issue", "detail"])
        w.writeheader()
        w.writerows(issues)

    split_docs = defaultdict(int)
    split_examples = {k: len(v) for k, v in sft.items()}
    processes_by_split: dict[str, set[str]] = defaultdict(set)
    for rec in gold_records:
        split_docs[rec["split"]] += 1
        processes_by_split[rec["split"]].add(rec["process_hash"])

    overlap = set()
    splits = ["train", "val", "test"]
    for i in range(len(splits)):
        for j in range(i+1, len(splits)):
            overlap |= processes_by_split[splits[i]] & processes_by_split[splits[j]]

    manifest = {
        "version": "3.1.0-local-xlsx-teacher-gold-slm",
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "review_xlsx": str(review),
        "validated_documents": len(gold_records),
        "gold_recortes": sum(len(r["gold"]) for r in gold_records),
        "documents_by_split": dict(split_docs),
        "sft_examples_by_split": split_examples,
        "process_overlap_between_splits": len(overlap),
        "seed": args.seed,
        "split_unit": "Processo",
        "issues": len(issues),
    }
    (out / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")

    print(json.dumps(manifest, ensure_ascii=False, indent=2))
    print("GOLD:", gold_path)
    print("SFT:", out / "sft_train.jsonl")
    print("Issues:", issues_path)


if __name__ == "__main__":
    main()
