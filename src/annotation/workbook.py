from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


DECISION_HEADERS = [
    "source_row",
    "Extração",
    "Processo",
    "Classe judicial",
    "Órgão julgador",
    "Polo Ativo",
    "Polo Passivo",
    "Decisão",
    "Matéria SAJ",
    "Ind.",
    "tags",
    "teacher_model",
    "pipeline_version",
    "generated_at_utc",
    "teacher_recortes_count",
    "status_documento",
    "observacao_documento",
    "revisor_documento",
    "revisado_em",
]

CANDIDATE_HEADERS = [
    "source_row",
    "Processo",
    "candidate_rank",
    "unit_ids",
    "section",
    "candidate_text",
    "teacher_category",
    "teacher_priority",
    "heuristic_score",
    "status_validacao",
    "texto_ajustado",
    "categoria_ajustada",
    "observacao",
    "revisor",
    "revisado_em",
]

ADDITION_HEADERS = [
    "source_row",
    "Processo",
    "gold_text_adicionado",
    "gold_category_adicionada",
    "observacao",
    "revisor",
    "revisado_em",
]

ALLOWED_CATEGORIES = [
    "resultado_julgamento",
    "ordem_determinacao",
    "intimacao_manifestacao",
    "tutela",
    "obrigacao",
    "prazo_cumprimento",
    "restituicao_pagamento",
    "honorarios_custas",
    "recurso_proximo_passo",
    "prescricao_decadencia",
    "reconhecimento_concordancia",
    "outro_acionavel",
]

DOC_STATUSES = ["VALIDADO_COMPLETO", "SEM_RECORTE", "PRECISA_REVISAO"]
CANDIDATE_STATUSES = ["APROVADO", "AJUSTADO", "REJEITADO", "SEM_RECORTE"]


def utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _style_header(ws) -> None:
    fill = PatternFill("solid", fgColor="D9EAF7")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def _set_widths(ws, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[col].width = width


def create_review_workbook(path: Path, manifest: dict[str, Any]) -> Workbook:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws_dec = wb.active
    ws_dec.title = "decisoes"
    ws_dec.append(DECISION_HEADERS)

    ws_can = wb.create_sheet("candidatos")
    ws_can.append(CANDIDATE_HEADERS)

    ws_add = wb.create_sheet("adicoes_gold")
    ws_add.append(ADDITION_HEADERS)

    ws_ins = wb.create_sheet("instrucoes")
    instructions = [
        ["PASSO", "INSTRUÇÃO"],
        [1, "Leia a decisão completa na aba decisoes."],
        [2, "Na aba candidatos, marque cada candidato como APROVADO, AJUSTADO ou REJEITADO."],
        [3, "Se o teacher deixou escapar um recorte, registre-o na aba adicoes_gold."],
        [4, "Só marque status_documento=VALIDADO_COMPLETO depois de revisar a decisão inteira."],
        [5, "Use SEM_RECORTE somente quando a decisão realmente não tiver recorte acionável."],
        [6, "Texto ajustado/adicionado deve ser literal da decisão. O builder de GOLD valida isso."],
        [7, "tags e Ind. são preservados como referência, mas NÃO foram enviados ao teacher."],
    ]
    for row in instructions:
        ws_ins.append(row)

    ws_man = wb.create_sheet("manifesto")
    ws_man.append(["chave", "valor"])
    for key, value in manifest.items():
        ws_man.append([key, str(value)])

    for ws in [ws_dec, ws_can, ws_add, ws_ins, ws_man]:
        _style_header(ws)
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    _set_widths(ws_dec, {
        "A": 12, "B": 18, "C": 26, "D": 24, "E": 24, "F": 24, "G": 24,
        "H": 90, "I": 24, "J": 26, "K": 22, "L": 30, "M": 20,
        "N": 25, "O": 14, "P": 22, "Q": 40, "R": 24, "S": 22,
    })
    _set_widths(ws_can, {
        "A": 12, "B": 26, "C": 12, "D": 18, "E": 18, "F": 85,
        "G": 30, "H": 16, "I": 16, "J": 20, "K": 85, "L": 30,
        "M": 40, "N": 24, "O": 22,
    })
    _set_widths(ws_add, {"A": 12, "B": 26, "C": 90, "D": 30, "E": 40, "F": 24, "G": 22})
    _set_widths(ws_ins, {"A": 10, "B": 110})
    _set_widths(ws_man, {"A": 35, "B": 70})

    dv_doc = DataValidation(type="list", formula1='"' + ','.join(DOC_STATUSES) + '"', allow_blank=True)
    ws_dec.add_data_validation(dv_doc)
    dv_doc.add("P2:P1048576")

    dv_status = DataValidation(type="list", formula1='"' + ','.join(CANDIDATE_STATUSES) + '"', allow_blank=True)
    ws_can.add_data_validation(dv_status)
    dv_status.add("J2:J1048576")

    dv_cat = DataValidation(type="list", formula1='"' + ','.join(ALLOWED_CATEGORIES) + '"', allow_blank=True)
    ws_can.add_data_validation(dv_cat)
    dv_cat.add("L2:L1048576")

    dv_add_cat = DataValidation(type="list", formula1='"' + ','.join(ALLOWED_CATEGORIES) + '"', allow_blank=True)
    ws_add.add_data_validation(dv_add_cat)
    dv_add_cat.add("D2:D1048576")

    wb.save(path)
    return wb


def load_or_create_review_workbook(path: Path, manifest: dict[str, Any], resume: bool) -> Workbook:
    if resume and path.exists():
        return load_workbook(path)
    return create_review_workbook(path, manifest)


def row_to_dict(headers: list[str], values: tuple[Any, ...]) -> dict[str, Any]:
    return {headers[i]: values[i] for i in range(min(len(headers), len(values)))}
