from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterator

from openpyxl import load_workbook


EXPECTED_COLUMNS = [
    "Extração",
    "Processo",
    "Classe judicial",
    "Órgão julgador",
    "Polo Ativo",
    "Polo Passivo",
    "Decisão",
    "Matéria SAJ",
    "Ind.",
    "tags",
]

REQUIRED_COLUMNS = ["Processo", "Decisão"]


def _norm_header(value: Any) -> str:
    text = str(value or "").strip()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = text.casefold()
    text = re.sub(r"[^a-z0-9]+", " ", text)
    return re.sub(r"\s+", " ", text).strip()


_CANONICAL_BY_NORMALIZED = {
    _norm_header(name): name for name in EXPECTED_COLUMNS
}
# Aliases tolerated without changing the canonical output contract.
_CANONICAL_BY_NORMALIZED.update({
    "extracao": "Extração",
    "processo": "Processo",
    "classe judicial": "Classe judicial",
    "orgao julgador": "Órgão julgador",
    "polo ativo": "Polo Ativo",
    "polo passivo": "Polo Passivo",
    "decisao": "Decisão",
    "materia saj": "Matéria SAJ",
    "ind": "Ind.",
    "indicador": "Ind.",
    "tags": "tags",
    "tag": "tags",
})


@dataclass(frozen=True)
class LocalWorkbookInfo:
    path: Path
    sheet: str
    headers: list[str]
    canonical_to_index: dict[str, int]
    max_row: int


def inspect_workbook(path: str | Path, sheet: str = "input") -> LocalWorkbookInfo:
    source = Path(path)
    if not source.exists():
        raise FileNotFoundError(f"Planilha de entrada não encontrada: {source}")
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError("Formato não suportado. Use .xlsx ou .xlsm.")

    wb = load_workbook(source, read_only=True, data_only=True)
    try:
        if sheet not in wb.sheetnames:
            raise ValueError(
                f"Aba '{sheet}' não encontrada. Abas disponíveis: {', '.join(wb.sheetnames)}"
            )
        ws = wb[sheet]
        first_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), tuple())
        headers = [str(v or "").strip() for v in first_row]
        canonical_to_index: dict[str, int] = {}
        for idx, raw in enumerate(headers):
            canonical = _CANONICAL_BY_NORMALIZED.get(_norm_header(raw))
            if canonical and canonical not in canonical_to_index:
                canonical_to_index[canonical] = idx

        missing = [name for name in REQUIRED_COLUMNS if name not in canonical_to_index]
        if missing:
            raise ValueError(
                "Colunas obrigatórias ausentes: " + ", ".join(missing)
                + ". Cabeçalhos encontrados: " + " | ".join(headers)
            )

        return LocalWorkbookInfo(
            path=source,
            sheet=sheet,
            headers=headers,
            canonical_to_index=canonical_to_index,
            max_row=ws.max_row,
        )
    finally:
        wb.close()


def iter_source_rows(
    path: str | Path,
    *,
    sheet: str = "input",
    start_excel_row: int = 2,
    limit: int = 0,
) -> Iterator[dict[str, Any]]:
    info = inspect_workbook(path, sheet)
    wb = load_workbook(info.path, read_only=True, data_only=True)
    ws = wb[sheet]
    yielded = 0
    try:
        for excel_row, values in enumerate(
            ws.iter_rows(min_row=max(2, start_excel_row), values_only=True),
            start=max(2, start_excel_row),
        ):
            if limit and yielded >= limit:
                break
            row: dict[str, Any] = {"source_row": excel_row}
            for canonical in EXPECTED_COLUMNS:
                idx = info.canonical_to_index.get(canonical)
                row[canonical] = values[idx] if idx is not None and idx < len(values) else ""
            yielded += 1
            yield row
    finally:
        wb.close()
