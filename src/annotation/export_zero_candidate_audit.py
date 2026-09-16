from __future__ import annotations

import argparse
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from src.annotation.local_xlsx import iter_source_rows


HEADERS = [
    "source_row",
    "Processo",
    "Classe judicial",
    "Órgão julgador",
    "Polo Ativo",
    "Polo Passivo",
    "Decisão",
    "Matéria SAJ",
    "Ind.",
    "tags",
    "motivo_auditoria",
    "tem_recorte_acionavel",
    "observacao_revisor",
]


def _zero_rows(review_xlsx: Path) -> set[int]:
    wb = load_workbook(review_xlsx, read_only=True, data_only=True)
    try:
        if "candidatos" not in wb.sheetnames:
            raise ValueError("A planilha de revisão não contém a aba 'candidatos'.")
        ws = wb["candidatos"]
        rows = ws.iter_rows(values_only=True)
        header = next(rows, None)
        if not header:
            return set()
        idx = {str(v or "").strip(): i for i, v in enumerate(header)}
        if "source_row" not in idx or "candidate_rank" not in idx:
            raise ValueError("A aba candidatos não contém source_row/candidate_rank.")
        out: set[int] = set()
        for values in rows:
            try:
                rank = int(values[idx["candidate_rank"]] or 0)
                source_row = int(values[idx["source_row"]])
            except (TypeError, ValueError):
                continue
            if rank == 0:
                out.add(source_row)
        return out
    finally:
        wb.close()


def export_audit(
    *,
    input_xlsx: Path,
    review_xlsx: Path,
    output_xlsx: Path,
    sheet: str = "input",
) -> int:
    target_rows = _zero_rows(review_xlsx)
    output_xlsx.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "zero_candidates"
    ws.append(HEADERS)

    fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    count = 0
    if target_rows:
        min_row, max_row = min(target_rows), max(target_rows)
        for row in iter_source_rows(
            input_xlsx,
            sheet=sheet,
            start_excel_row=min_row,
            limit=max_row - min_row + 1,
        ):
            source_row = int(row["source_row"])
            if source_row not in target_rows:
                continue
            ws.append([
                source_row,
                row.get("Processo", ""),
                row.get("Classe judicial", ""),
                row.get("Órgão julgador", ""),
                row.get("Polo Ativo", ""),
                row.get("Polo Passivo", ""),
                row.get("Decisão", ""),
                row.get("Matéria SAJ", ""),
                row.get("Ind.", ""),
                row.get("tags", ""),
                "teacher não selecionou candidato; revisar a decisão completa",
                "",
                "",
            ])
            count += 1

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    widths = {
        "A": 12, "B": 26, "C": 24, "D": 24, "E": 24, "F": 24,
        "G": 100, "H": 24, "I": 24, "J": 24, "K": 48, "L": 24, "M": 48,
    }
    for col, width in widths.items():
        ws.column_dimensions[col].width = width
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(output_xlsx)
    return count


def main() -> None:
    p = argparse.ArgumentParser(
        description="Exporta decisões que ficaram com candidate_rank=0 para auditoria humana local."
    )
    p.add_argument("--input-xlsx", required=True)
    p.add_argument("--review-xlsx", required=True)
    p.add_argument("--sheet", default="input")
    p.add_argument("--output", required=True)
    args = p.parse_args()

    count = export_audit(
        input_xlsx=Path(args.input_xlsx),
        review_xlsx=Path(args.review_xlsx),
        output_xlsx=Path(args.output),
        sheet=args.sheet,
    )
    print("=== DIDE1 / AUDITORIA DE ZERO CANDIDATOS ===")
    print("decisões exportadas:", count)
    print("arquivo:", args.output)
    print("IMPORTANTE: leia a decisão completa e marque se realmente não há recorte acionável.")


if __name__ == "__main__":
    main()
