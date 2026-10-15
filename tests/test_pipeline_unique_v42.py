from pathlib import Path
from openpyxl import Workbook, load_workbook

from dide1.pipeline import executar
from dide1.teacher import ResultadoModelo


def _criar_xlsx(path: Path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Processo", "Classe", "Origem", "Órgão Julgador", "Expediente", "Assunto", "Seção/Subseção", "Decisão Judicial", "Núcleo"])
    d1 = "DESPACHO\n\nIntime-se a parte A para manifestação no prazo de 10 dias."
    d2 = "DESPACHO\n\nIntime-se a parte B para manifestação no prazo de 15 dias."
    d3 = "DESPACHO\n\nIntime-se a parte C para manifestação no prazo de 20 dias."
    ws.append(["P1", "CumSenFaz", "PJE-AL", "1ª Vara Federal AL", "Intimação (1)", "Tema", "AL / Maceió", d1, "NTDC"])
    ws.append(["P1", "CumSenFaz", "PJE-AL", "1ª Vara Federal AL", "Intimação (2)", "Tema", "AL / Maceió", d1, "NTDC"])
    ws.append(["P2", "CumSenFaz", "PJE-AL", "1ª Vara Federal AL", "Intimação (3)", "Tema", "AL / Maceió", d2, "NTDC"])
    ws.append(["P3", "CumSenFaz", "PJE-AL", "1ª Vara Federal AL", "Intimação (4)", "Tema", "AL / Maceió", d3, "NTDC"])
    wb.save(path)


def test_limite_unicos_gera_quantidade_exata_de_decisoes(tmp_path, monkeypatch):
    entrada = tmp_path / "in.xlsx"
    saida = tmp_path / "out"
    _criar_xlsx(entrada)

    def chamada_falsa(registro, unidades, roteamento, regra_encontrada, maximo_candidatos=3):
        acionavel = [u for u in unidades if "Intime-se" in u.texto][0]
        return ResultadoModelo(
            selecionados=[{"unit_ids": [acionavel.id_unidade], "category": "intimacao_manifestacao", "priority": 1, "reason": "ok"}],
            situacao="OK", tentativas=1, modo_resposta="json_schema",
        )

    monkeypatch.setattr("dide1.pipeline.chamar_modelo", chamada_falsa)
    manifesto = executar(str(entrada), str(saida), modo_modelo="llama", limite_unicos=2)
    assert manifesto["unique_decisions"] == 2
    assert manifesto["scanned"] == 3
    assert manifesto["processed"] == 3
    assert manifesto["cache_hits"] == 1

    wb = load_workbook(saida / "review.xlsx")
    assert wb["revisao"].max_row == 3  # cabeçalho + 2 decisões únicas
