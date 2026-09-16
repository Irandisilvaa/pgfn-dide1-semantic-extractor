from pathlib import Path
from openpyxl import load_workbook
from src.annotation.workbook import create_review_workbook


def test_review_workbook_has_required_sheets(tmp_path: Path):
    path = tmp_path / "review.xlsx"
    create_review_workbook(path, {"pipeline_version": "test"})
    wb = load_workbook(path)
    assert {"decisoes", "candidatos", "adicoes_gold", "instrucoes", "manifesto"}.issubset(wb.sheetnames)
