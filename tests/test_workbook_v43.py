from pathlib import Path
from openpyxl import load_workbook

from dide1.workbook import escrever_planilha_revisao, construir_gold, reconstruir_revisao


def _decisao(linha_origem, decision_hash, expediente, cache=False):
    return {
        "source_row": linha_origem,
        "processo": "P1",
        "classe": "CumSenFaz",
        "origem": "PJE-AL",
        "orgao_julgador": "1ª Vara Federal AL",
        "data": "2026-01-01",
        "expediente": expediente,
        "prazo": 15,
        "assunto": "Tema",
        "secao_subsecao": "AL / Maceió",
        "link": f"https://example/{linha_origem}",
        "decisao": "DESPACHO\n\nIntime-se a parte em 15 dias.",
        "nucleo": "NTDC",
        "rule_id": "J19",
        "rule_label": "ato ordinatório",
        "route_key": "cumsenfaz|al_maceio|1a_vara",
        "specialist": "global_teacher",
        "teacher_status": "OK",
        "teacher_attempts": 1,
        "teacher_response_mode": "json_schema",
        "decision_sha256": decision_hash,
        "cache_hit": cache,
    }


def _candidato(linha_origem, decision_hash):
    return {
        "source_row": linha_origem,
        "Processo": "P1",
        "candidate_rank": 1,
        "candidate_text": "Intime-se a parte em 15 dias.",
        "teacher_category": "intimacao_manifestacao",
        "teacher_priority": 1,
        "teacher_reason": "ordem clara",
        "decision_sha256": decision_hash,
    }


def test_v43_revisao_compacta_oculta_dados_tecnicos(tmp_path: Path):
    path = tmp_path / "review.xlsx"
    escrever_planilha_revisao(path, [_decisao(2, "abc", "Intimação (1)"), _decisao(3, "abc", "Intimação (2)", True)], [_candidato(2, "abc"), _candidato(3, "abc")])
    wb = load_workbook(path)
    assert wb.sheetnames[0] == "revisao"
    assert wb["dados_tecnicos"].sheet_state == "hidden"
    assert wb["candidatos_tecnicos"].sheet_state == "hidden"
    ws = wb["revisao"]
    headers = {c.value: c.column for c in ws[1]}
    assert "Recorte 1" in headers
    assert "Validação 1" in headers
    assert "candidate_1_reason" not in headers
    assert ws.cell(2, headers["Validação 1"]).value == "PENDENTE"
    assert ws.column_dimensions["A"].hidden is True


def test_v43_gold_usa_valores_visiveis_editados(tmp_path: Path):
    revisao = tmp_path / "review.xlsx"
    saida = tmp_path / "gold.jsonl"
    escrever_planilha_revisao(revisao, [_decisao(2, "abc", "Intimação (1)")], [_candidato(2, "abc")])
    wb = load_workbook(revisao)
    ws = wb["revisao"]
    h = {c.value: c.column for c in ws[1]}
    ws.cell(2, h["Recorte 1"]).value = "Intime-se a parte em 10 dias."
    ws.cell(2, h["Categoria 1"]).value = "intimacao_manifestacao"
    ws.cell(2, h["Validação 1"]).value = "AJUSTADO"
    ws.cell(2, h["Status da decisão"]).value = "VALIDADO_COMPLETO"
    wb.save(revisao)
    resultado = construir_gold(str(revisao), str(saida))
    assert resultado["validated_documents"] == 1
    assert resultado["blocked_documents"] == 0
    texto = saida.read_text(encoding="utf-8")
    assert "10 dias" in texto
    assert '"validation": "AJUSTADO"' in texto


def test_v43_bloqueia_documento_validado_incompleto(tmp_path: Path):
    revisao = tmp_path / "review.xlsx"
    saida = tmp_path / "gold.jsonl"
    escrever_planilha_revisao(revisao, [_decisao(2, "abc", "Intimação (1)")], [_candidato(2, "abc")])
    wb = load_workbook(revisao)
    ws = wb["revisao"]
    h = {c.value: c.column for c in ws[1]}
    ws.cell(2, h["Status da decisão"]).value = "VALIDADO_COMPLETO"
    wb.save(revisao)
    resultado = construir_gold(str(revisao), str(saida))
    assert resultado["validated_documents"] == 0
    assert resultado["blocked_documents"] == 1
    issues = (tmp_path / "gold_build_issues.csv").read_text(encoding="utf-8-sig")
    assert "ainda não revisado" in issues


def test_reconstruir_revisao_nao_exige_qwen(tmp_path: Path):
    old = tmp_path / "old.xlsx"
    new = tmp_path / "new.xlsx"
    # Simula v4.3 para garantir caminho de rebuild; a função também aceita v4.2.
    escrever_planilha_revisao(old, [_decisao(2, "abc", "Intimação (1)")], [_candidato(2, "abc")])
    resultado = reconstruir_revisao(str(old), str(new))
    assert resultado["qwen_rerun"] is False
    assert new.exists()
