from pathlib import Path
from openpyxl import load_workbook

from dide1.workbook import escrever_planilha_revisao, construir_gold


def _decisao(linha_origem, decision_hash, expediente, cache=False):
    return {
        "source_row": linha_origem, "processo": "P1", "classe": "CumSenFaz", "origem": "PJE-AL",
        "orgao_julgador": "1ª Vara Federal AL", "data": "2026-01-01", "expediente": expediente,
        "prazo": 15, "assunto": "Tema", "secao_subsecao": "AL / Maceió",
        "link": f"https://example/{linha_origem}", "decisao": "DESPACHO\n\nIntime-se a parte em 15 dias.",
        "nucleo": "NTDC", "rule_id": "J19", "rule_label": "ato ordinatório",
        "route_key": "cumsenfaz|al_maceio|1a_vara", "specialist": "global_teacher",
        "teacher_status": "OK", "teacher_attempts": 1, "teacher_response_mode": "json_schema",
        "decision_sha256": decision_hash, "cache_hit": cache,
    }


def _candidato(linha_origem, decision_hash):
    return {
        "source_row": linha_origem, "Processo": "P1", "candidate_rank": 1,
        "candidate_text": "Intime-se a parte em 15 dias.",
        "teacher_category": "intimacao_manifestacao", "teacher_priority": 1,
        "teacher_reason": "ordem clara", "decision_sha256": decision_hash,
    }


def test_planilha_revisao_consolida_decisoes_identicas(tmp_path: Path):
    path = tmp_path / "review.xlsx"
    decisions = [_decisao(2, "abc", "Intimação (1)"), _decisao(3, "abc", "Intimação (2)", True)]
    candidates = [_candidato(2, "abc"), _candidato(3, "abc")]
    escrever_planilha_revisao(path, decisions, candidates)
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames[0] == "revisao"
    ws = wb["revisao"]
    assert ws.max_row == 2
    headers = {c.value: c.column for c in ws[1]}
    assert ws.cell(2, headers["Recorte 1"]).value == "Intime-se a parte em 15 dias."
    assert ws.cell(2, headers["Validação 1"]).value == "PENDENTE"
    tech = wb["dados_tecnicos"]
    th = {c.value: c.column for c in tech[1]}
    assert tech.cell(2, th["ocorrencias_na_base"]).value == 2


def test_construcao_gold_le_revisao_compacta(tmp_path: Path):
    caminho_revisao = tmp_path / "review.xlsx"
    saida = tmp_path / "gold.jsonl"
    escrever_planilha_revisao(caminho_revisao, [_decisao(2, "abc", "Intimação (1)")], [_candidato(2, "abc")])
    wb = load_workbook(caminho_revisao)
    ws = wb["revisao"]
    headers = {c.value: c.column for c in ws[1]}
    ws.cell(2, headers["Validação 1"]).value = "APROVADO"
    ws.cell(2, headers["Recorte ausente 1"]).value = "Arquivem-se os autos."
    ws.cell(2, headers["Categoria ausente 1"]).value = "recurso_proximo_passo"
    ws.cell(2, headers["Status da decisão"]).value = "VALIDADO_COMPLETO"
    wb.save(caminho_revisao)
    resultado = construir_gold(str(caminho_revisao), str(saida))
    assert resultado["validated_documents"] == 1
    conteudo = saida.read_text(encoding="utf-8")
    assert "Intime-se a parte em 15 dias." in conteudo
    assert "Arquivem-se os autos." in conteudo
