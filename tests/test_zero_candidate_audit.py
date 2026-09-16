from pathlib import Path

from openpyxl import Workbook, load_workbook

from src.annotation.export_zero_candidate_audit import export_audit


def test_export_zero_candidate_audit(tmp_path: Path):
    source = tmp_path / "source.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "input"
    ws.append([
        "Extração", "Processo", "Classe judicial", "Órgão julgador",
        "Polo Ativo", "Polo Passivo", "Decisão", "Matéria SAJ", "Ind.", "tags"
    ])
    ws.append(["", "P1", "MS", "Vara", "A", "B", "Decisão um", "", "", ""])
    ws.append(["", "P2", "MS", "Vara", "A", "B", "Decisão dois", "", "", ""])
    wb.save(source)

    review = tmp_path / "review.xlsx"
    wb2 = Workbook()
    ws2 = wb2.active
    ws2.title = "candidatos"
    ws2.append(["source_row", "candidate_rank"])
    ws2.append([2, 1])
    ws2.append([3, 0])
    wb2.save(review)

    output = tmp_path / "audit.xlsx"
    count = export_audit(
        input_xlsx=source,
        review_xlsx=review,
        output_xlsx=output,
        sheet="input",
    )
    assert count == 1

    wb3 = load_workbook(output, read_only=True, data_only=True)
    try:
        ws3 = wb3["zero_candidates"]
        rows = list(ws3.iter_rows(values_only=True))
        assert rows[1][0] == 3
        assert rows[1][1] == "P2"
        assert rows[1][6] == "Decisão dois"
    finally:
        wb3.close()
